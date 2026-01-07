from __future__ import annotations

import os
import re
import sqlite3
from dataclasses import dataclass

from docx import Document
from docx.shared import RGBColor

INPUT_DOC = "docs/references/BRHC1914-final.docx"
DB_PATH = "assets/databases/brhc.db"

CHAPTER_RE = re.compile(r"^(?:\[Ch\]\s*)?Chapter\s+(\d+)", re.IGNORECASE)
KNOWN_MARKERS = {"[I]", "[S]", "[Ch]", "[N]", "[P]", "[R]", "[T]"}


@dataclass(frozen=True)
class ImageBlock:
    block_id: int
    section_title: str | None
    chapter_title: str | None
    block_order: int
    block_type: str
    raw_text: str | None
    normalized_text: str | None
    table_json: str | None
    image_blob_id: int | None


def _is_blue(run) -> bool:
    color = run.font.color
    return color is not None and color.rgb == RGBColor(0, 0, 255)


def _escape_html(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _render_table(lines: list[str]) -> str:
    rows: list[list[str]] = []
    for line in lines:
        if "|" in line:
            rows.append([cell for cell in line.split("|")])
        elif "\t" in line:
            rows.append([cell for cell in line.split("\t")])
        else:
            return "<pre>" + _escape_html("\n".join(lines)) + "</pre>"
    col_count = max(len(row) for row in rows)
    if any(len(row) != col_count for row in rows):
        return "<pre>" + _escape_html("\n".join(lines)) + "</pre>"
    html_rows = []
    for row in rows:
        cells = "".join(f"<td>{_escape_html(cell)}</td>" for cell in row)
        html_rows.append(f"<tr>{cells}</tr>")
    return "<table>" + "".join(html_rows) + "</table>"


def _load_image_blocks(cur: sqlite3.Cursor) -> dict[tuple, list[ImageBlock]]:
    cur.execute("SELECT DISTINCT block_id FROM brhc_image_block_map")
    ids = [row[0] for row in cur.fetchall()]
    if not ids:
        return {}
    placeholders = ",".join(["?"] * len(ids))
    cur.execute(
        f"""
        SELECT block_id, section_title, chapter_title, block_order, block_type,
               raw_text, normalized_text, table_json, image_blob_id
        FROM doc_blocks
        WHERE block_id IN ({placeholders})
        """,
        ids,
    )
    image_blocks: dict[tuple, list[ImageBlock]] = {}
    for row in cur.fetchall():
        block = ImageBlock(
            block_id=row[0],
            section_title=row[1],
            chapter_title=row[2],
            block_order=row[3],
            block_type=row[4],
            raw_text=row[5],
            normalized_text=row[6],
            table_json=row[7],
            image_blob_id=row[8],
        )
        key = (block.block_type, block.raw_text, block.section_title, block.chapter_title)
        image_blocks.setdefault(key, []).append(block)
    return image_blocks


def _next_block_id(cur: sqlite3.Cursor) -> int:
    cur.execute("SELECT COALESCE(MAX(block_id), 0) FROM doc_blocks")
    return int(cur.fetchone()[0]) + 1


def _insert_section(cur: sqlite3.Cursor, title: str, order_index: int) -> int:
    cur.execute(
        "INSERT INTO brhc_sections (section_title, order_index) VALUES (?, ?)",
        (title, order_index),
    )
    return int(cur.lastrowid)


def _insert_chapter(
    cur: sqlite3.Cursor,
    section_id: int | None,
    title: str,
    order_index: int,
) -> int:
    cur.execute(
        "INSERT INTO brhc_chapters (section_id, chapter_title, order_index) VALUES (?, ?, ?)",
        (section_id, title, order_index),
    )
    return int(cur.lastrowid)


def _insert_doc_block(
    cur: sqlite3.Cursor,
    block_id: int,
    section_title: str | None,
    chapter_title: str | None,
    block_order: int,
    block_type: str,
    raw_text: str | None,
    normalized_text: str | None,
    table_json: str | None = None,
    image_blob_id: int | None = None,
) -> None:
    cur.execute(
        """
        INSERT INTO doc_blocks (
            block_id, section_title, chapter_title, block_order, block_type,
            raw_text, normalized_text, table_json, image_blob_id
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            block_id,
            section_title,
            chapter_title,
            block_order,
            block_type,
            raw_text,
            normalized_text,
            table_json,
            image_blob_id,
        ),
    )


def _insert_question(
    cur: sqlite3.Cursor,
    block_id: int,
    question_number: int,
    question_text: str,
    section_title: str | None,
    chapter_title: str | None,
) -> None:
    cur.execute(
        """
        INSERT INTO d_questions (
            block_id, question_number, question_text, section_title, chapter_title
        ) VALUES (?, ?, ?, ?, ?)
        """,
        (block_id, question_number, question_text, section_title, chapter_title),
    )


def _insert_question_row(
    cur: sqlite3.Cursor,
    section_title: str | None,
    chapter_title: str | None,
    chapter_number: int | None,
    question_number: int,
    question_text: str,
) -> None:
    cur.execute(
        """
        INSERT INTO questions (
            section_title, chapter_title, chapter_number, content_type,
            question_number, question_text, answer_text
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            section_title,
            chapter_title,
            chapter_number,
            "question",
            question_number,
            question_text,
            None,
        ),
    )


def _delete_content_tables(cur: sqlite3.Cursor) -> None:
    cur.execute("DELETE FROM brhc_sections")
    cur.execute("DELETE FROM brhc_chapters")
    cur.execute("DELETE FROM doc_blocks")
    cur.execute("DELETE FROM d_questions")
    cur.execute("DELETE FROM questions")


def _validate(cur: sqlite3.Cursor) -> list[str]:
    issues: list[str] = []
    cur.execute(
        """
        SELECT COUNT(*)
        FROM brhc_chapters c
        LEFT JOIN brhc_sections s ON s.section_id = c.section_id
        WHERE c.section_id IS NOT NULL AND s.section_id IS NULL
        """
    )
    if cur.fetchone()[0] != 0:
        issues.append("Orphan chapters detected.")

    cur.execute(
        """
        SELECT COUNT(*)
        FROM doc_blocks
        WHERE chapter_title IS NULL
          AND block_type NOT IN ('intro_heading', 'intro_paragraph', 'section')
        """
    )
    if cur.fetchone()[0] != 0:
        issues.append("Orphan blocks detected (missing chapter_title).")

    cur.execute(
        """
        SELECT chapter_title, COUNT(*) AS cnt, MAX(question_number) AS max_q
        FROM d_questions
        GROUP BY chapter_title
        HAVING cnt != max_q OR MIN(question_number) != 1
        """
    )
    if cur.fetchall():
        issues.append("Question numbering not sequential per chapter.")

    return issues


def main() -> None:
    if not os.path.exists(INPUT_DOC):
        raise FileNotFoundError(f"Missing input document: {INPUT_DOC}")
    if not os.path.exists(DB_PATH):
        raise FileNotFoundError(f"Missing database: {DB_PATH}")

    conn = sqlite3.connect(DB_PATH, timeout=30)
    cur = conn.cursor()
    cur.execute("PRAGMA foreign_keys = OFF")
    cur.execute("PRAGMA busy_timeout = 30000")
    cur.execute("BEGIN")

    image_block_map = _load_image_blocks(cur)
    image_blocks_by_chapter: dict[str | None, list[ImageBlock]] = {}
    image_block_ids = {b.block_id for blocks in image_block_map.values() for b in blocks}
    for blocks in image_block_map.values():
        for block in blocks:
            image_blocks_by_chapter.setdefault(block.chapter_title, []).append(block)
    next_block_id = _next_block_id(cur)

    _delete_content_tables(cur)

    doc = Document(INPUT_DOC)

    section_order = 0
    chapter_order = 0
    current_section_title = None
    current_section_id = None
    current_chapter_title = None
    current_chapter_num = None
    block_order = 0
    question_counter = 0

    sections_total = 0
    chapters_processed = 0
    questions_total = 0
    notes_total = 0
    poetry_total = 0
    responsive_total = 0
    tables_total = 0

    anomalies: list[str] = []
    used_block_ids: set[int] = set()

    intro_active = True
    aggregated_type: str | None = None
    aggregated_lines: list[str] = []

    current_mode: str | None = None
    current_buffer: list[str] = []
    current_has_text = False

    def allocate_block_id(
        block_type: str, raw_text: str | None, block_order_value: int
    ) -> int:
        nonlocal next_block_id
        key = (block_type, raw_text, current_section_title, current_chapter_title)
        candidates = image_block_map.get(key)
        if candidates:
            while candidates:
                block = candidates.pop(0)
                if block.block_id in used_block_ids:
                    continue
                used_block_ids.add(block.block_id)
                if block in image_blocks_by_chapter.get(block.chapter_title, []):
                    image_blocks_by_chapter[block.chapter_title].remove(block)
                return block.block_id
        if current_chapter_title in image_blocks_by_chapter:
            blocks = [
                b
                for b in image_blocks_by_chapter[current_chapter_title]
                if b.block_id not in used_block_ids
            ]
            if blocks:
                same_type = [b for b in blocks if b.block_type == block_type]
                pool = same_type if same_type else blocks
                pick = min(
                    pool, key=lambda b: abs(b.block_order - block_order_value)
                )
                image_blocks_by_chapter[current_chapter_title].remove(pick)
                used_block_ids.add(pick.block_id)
                anomalies.append(
                    f"Reassigned image block_id={pick.block_id} to {block_type} near order {block_order_value}"
                )
                return pick.block_id
        block_id = next_block_id
        next_block_id += 1
        return block_id

    def flush_aggregate() -> None:
        nonlocal aggregated_type, aggregated_lines, block_order
        nonlocal poetry_total, responsive_total, tables_total
        if aggregated_type is None or not aggregated_lines:
            aggregated_type = None
            aggregated_lines = []
            return
        block_order += 1
        raw_text = "\n".join(aggregated_lines)
        block_id = allocate_block_id(aggregated_type, raw_text, block_order)
        payload = raw_text
        if aggregated_type == "table":
            payload = _render_table(aggregated_lines)
        _insert_doc_block(
            cur,
            block_id,
            current_section_title,
            current_chapter_title,
            block_order,
            aggregated_type,
            raw_text,
            payload,
        )
        if aggregated_type == "poetry":
            poetry_total += 1
        elif aggregated_type == "responsive":
            responsive_total += 1
        elif aggregated_type == "table":
            tables_total += 1
        aggregated_type = None
        aggregated_lines = []

    def flush_qa_block() -> None:
        nonlocal current_mode, current_buffer, current_has_text, block_order
        nonlocal question_counter, questions_total
        if current_mode is None or not current_buffer:
            current_mode = None
            current_buffer = []
            current_has_text = False
            return
        text = "".join(current_buffer)
        block_order += 1
        block_type = "question" if current_mode == "question" else "answer"
        block_id = allocate_block_id(block_type, text, block_order)
        _insert_doc_block(
            cur,
            block_id,
            current_section_title,
            current_chapter_title,
            block_order,
            block_type,
            text,
            text,
        )
        if current_mode == "question":
            question_counter += 1
            questions_total += 1
            _insert_question(
                cur,
                block_id,
                question_counter,
                text,
                current_section_title,
                current_chapter_title,
            )
            _insert_question_row(
                cur,
                current_section_title,
                current_chapter_title,
                current_chapter_num,
                question_counter,
                text,
            )
        else:
            pass
        current_mode = None
        current_buffer = []
        current_has_text = False

    def start_mode(mode: str) -> None:
        nonlocal current_mode
        if current_mode is None:
            current_mode = mode
            return
        if current_mode != mode:
            flush_qa_block()
            current_mode = mode

    for para in doc.paragraphs:
        raw_text = para.text
        stripped = raw_text.strip()

        if aggregated_type is not None:
            boundary = (
                stripped.startswith("[S]")
                or stripped.startswith("[Ch]")
                or stripped.startswith("[N]")
                or stripped.startswith("[R]")
                or stripped.startswith("[T]")
                or CHAPTER_RE.match(stripped)
            )
            if boundary:
                flush_aggregate()
            else:
                aggregated_lines.append(raw_text)
                continue

        if not stripped:
            if current_mode is not None and current_has_text:
                current_buffer.append("\n")
            continue

        if stripped.startswith("[") and "]" in stripped:
            marker = stripped[: stripped.index("]") + 1]
            if marker not in KNOWN_MARKERS:
                conn.rollback()
                raise RuntimeError(f"Unknown structural marker: {marker}")

        ch_match = CHAPTER_RE.match(stripped)
        if ch_match:
            flush_aggregate()
            flush_qa_block()
            current_chapter_num = int(ch_match.group(1))
            current_chapter_title = stripped
            question_counter = 0
            block_order = 0
            chapter_order += 1
            chapters_processed += 1
            intro_active = False
            _insert_chapter(cur, current_section_id, current_chapter_title, chapter_order)
            block_order += 1
            block_id = allocate_block_id("chapter", raw_text, block_order)
            _insert_doc_block(
                cur,
                block_id,
                current_section_title,
                current_chapter_title,
                block_order,
                "chapter",
                raw_text,
                raw_text,
            )
            continue

        if stripped.startswith("[S]"):
            flush_aggregate()
            flush_qa_block()
            current_section_title = stripped
            section_order += 1
            sections_total += 1
            current_section_id = _insert_section(cur, current_section_title, section_order)
            block_order += 1
            block_id = allocate_block_id("section", raw_text, block_order)
            _insert_doc_block(
                cur,
                block_id,
                current_section_title,
                current_chapter_title,
                block_order,
                "section",
                raw_text,
                raw_text,
            )
            intro_active = False
            continue

        if intro_active:
            if stripped.startswith("[I]"):
                flush_qa_block()
                block_order += 1
                block_id = allocate_block_id("intro", raw_text, block_order)
                content = raw_text
                if all(run.bold for run in para.runs if run.text.strip()):
                    content = f"<strong>{_escape_html(raw_text)}</strong>"
                    block_type = "intro_heading"
                else:
                    block_type = "intro_paragraph"
                _insert_doc_block(
                    cur,
                    block_id,
                    current_section_title,
                    current_chapter_title,
                    block_order,
                    block_type,
                    raw_text,
                    content,
                )
            continue

        if current_chapter_title is None:
            if any(_is_blue(run) for run in para.runs):
                anomalies.append(f"Orphan blue question before chapter: {stripped[:80]}")
            continue

        if stripped.startswith("[P]"):
            flush_qa_block()
            aggregated_type = "poetry"
            aggregated_lines = [raw_text]
            continue

        if stripped.startswith("[R]"):
            flush_qa_block()
            aggregated_type = "responsive"
            aggregated_lines = [raw_text]
            continue

        if stripped.startswith("[T]"):
            flush_qa_block()
            aggregated_type = "table"
            aggregated_lines = [raw_text]
            continue

        if stripped.startswith("[N]"):
            flush_qa_block()
            block_order += 1
            block_id = allocate_block_id("note", raw_text, block_order)
            _insert_doc_block(
                cur,
                block_id,
                current_section_title,
                current_chapter_title,
                block_order,
                "note",
                raw_text,
                raw_text,
            )
            notes_total += 1
            continue

        runs = [run for run in para.runs if run.text]
        if runs:
            first_mode = "question" if _is_blue(runs[0]) else "answer"
            if current_mode is not None and current_has_text:
                if current_mode == first_mode:
                    current_buffer.append("\n")
                else:
                    flush_qa_block()
        for run in runs:
            if not run.text:
                continue
            mode = "question" if _is_blue(run) else "answer"
            start_mode(mode)
            current_buffer.append(run.text)
            current_has_text = True

    flush_aggregate()
    flush_qa_block()
    if image_block_ids - used_block_ids:
        conn.rollback()
        missing = sorted(image_block_ids - used_block_ids)
        raise RuntimeError(
            "Unmatched image blocks; aborting to preserve references: "
            + ",".join(str(b) for b in missing)
        )

    issues = _validate(cur)
    if issues:
        conn.rollback()
        raise RuntimeError("Validation failed: " + "; ".join(issues))

    cur.execute("PRAGMA foreign_keys = ON")
    conn.commit()

    print("Import complete.")
    print(f"Sections imported: {sections_total}")
    print(f"Chapters imported: {chapters_processed}")
    print(f"Questions renumbered: {questions_total}")
    print(f"Notes imported: {notes_total}")
    print(f"Poetry blocks imported: {poetry_total}")
    print(f"Responsive readings imported: {responsive_total}")
    print(f"Tables imported: {tables_total}")
    print("Validation passed.")
    if anomalies:
        print("Anomalies:")
        for entry in anomalies:
            print(f"- {entry}")

    conn.close()


if __name__ == "__main__":
    main()
